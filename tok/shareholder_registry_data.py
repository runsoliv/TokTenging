"""Workbook adapter and searchable data model for the Hluthafaskrá service.

The source workbook is intentionally irregular.  This module treats its three
tab-selected sheets as source observations rather than flattening them into a
single lossy table.  Every populated cell in the meaningful bounds is retained
with its sheet, section, Excel coordinate, cached formula result, formula text,
and sensitivity classification.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field, replace
from pathlib import Path
import re
import threading
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


TARGET_SHEETS = (
    "Hlutafjármiðar 2025",
    "Hluthafaskrá",
    "Áreiðanleikakönnun allir",
)

# Explicit bounds avoid stale print areas and the corrupt XFD/bottom-row ranges
# present on unrelated workbook sheets.
SHEET_BOUNDS = {
    "Hlutafjármiðar 2025": (199, 18),  # A:R
    "Hluthafaskrá": (253, 19),         # A:S
    "Áreiðanleikakönnun allir": (231, 20),  # A:T
}

SHEET_PRIORITY = {
    "Hlutafjármiðar 2025": 0,
    "Áreiðanleikakönnun allir": 1,
    "Hluthafaskrá": 2,
}


class RegistryWorkbookError(ValueError):
    """Raised when a workbook does not contain the intended registry sheets."""


@dataclass(frozen=True)
class SourceCell:
    sheet: str
    section: str
    excel_row: int
    column: str
    column_index: int
    header: str
    original_value: object
    formula: str | None
    cached_value: object
    value: object
    sensitive: bool = False
    hyperlink: str | None = None

    @property
    def coordinate(self) -> str:
        return f"{self.column}{self.excel_row}"

    def display_value(self, reveal_sensitive: bool = False) -> str:
        if self.sensitive and not reveal_sensitive:
            return "••••••••"
        value = self.value
        if value is None and self.formula:
            value = self.formula
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()


@dataclass(frozen=True)
class SourceRow:
    sheet: str
    section: str
    excel_row: int
    cells: tuple[SourceCell, ...]
    record_type: str = "raw"
    layout_warning: str = ""


@dataclass(frozen=True)
class RelatedParty:
    role: str
    name: str
    kennitala: str
    share: object
    source_sheet: str
    excel_row: int
    warning: str = ""
    percentage: object = None


@dataclass(frozen=True)
class RegistryEntity:
    key: str
    kind: str
    kennitala: str
    name: str
    aliases: tuple[str, ...]
    emails: tuple[str, ...]
    source_rows: tuple[SourceRow, ...]
    related_parties: tuple[RelatedParty, ...]
    detail_rows: tuple[SourceRow, ...]
    warnings: tuple[str, ...]
    phones: tuple[str, ...] = ()
    statuses: tuple[str, ...] = ()
    reverse_ownerships: tuple[RelatedParty, ...] = ()

    @property
    def canonical_name(self) -> str:
        return self.name

    @property
    def normalized_kennitala(self) -> str:
        return self.kennitala


@dataclass(frozen=True)
class SearchResult:
    entity: RegistryEntity
    score: int
    matched_on: str
    matched_value: str


@dataclass(frozen=True)
class RegistryStats:
    entity_count: int
    company_count: int
    individual_count: int
    source_row_count: int
    source_cell_count: int
    sensitive_cell_count: int
    sheet_counts: dict[str, int]
    formula_cell_count: int = 0


@dataclass(frozen=True)
class _SectionSpec:
    sheet: str
    section: str
    start_row: int
    end_row: int
    header_row: int | None = None
    record_type: str = "raw"
    entity_kind: str = ""
    id_column: int | None = None
    name_column: int | None = None
    custom_headers: tuple[tuple[int, str], ...] = ()
    owner_pairs: tuple[tuple[int, int | None], ...] = ()
    warning: str = ""

    @property
    def is_entity(self) -> bool:
        return bool(self.entity_kind and self.name_column)


def _pairs(*values):
    return tuple(values)


SECTION_SPECS = (
    _SectionSpec(
        "Hlutafjármiðar 2025", "Hlutafjármiðar 2025", 4, 187, 3,
        "company", "company", 2, 3, owner_pairs=_pairs((6, 7), (8, 9), (10, 11), (12, 13)),
    ),
    _SectionSpec(
        "Hlutafjármiðar 2025", "Sundurliðun Búferla", 191, 191, None,
        "company", "company", 2, 3,
        custom_headers=_pairs((1, "Sundurliðun"), (2, "Kennitala"), (3, "Nafn"), (5, "Hlutafé")),
        warning="Í dálkum I–K í þessari línu er sérstakur útreikningur án fyrirsagnar.",
    ),
    _SectionSpec(
        "Hlutafjármiðar 2025", "Nánari sundurliðun Búferla", 193, 199, None,
        "detail", custom_headers=_pairs((1, "Eigandi"), (3, "Kennitala eiganda / land"), (5, "Hlutur"), (6, "Hlutfall")),
    ),

    _SectionSpec(
        "Hluthafaskrá", "A — Hlutafélög í áreiðanleikakönnun", 4, 97, 3,
        "company", "company", 2, 3, owner_pairs=_pairs((5, 6), (7, 8), (9, 10), (11, 12)),
    ),
    _SectionSpec(
        "Hluthafaskrá", "B — Nánari sundurliðun", 103, 103, None,
        "company", "company", 3, 4,
        custom_headers=_pairs((2, "Sundurliðun"), (3, "Kennitala"), (4, "Nafn"), (5, "Hlutafé")),
    ),
    _SectionSpec(
        "Hluthafaskrá", "B — Nánari sundurliðun Búferla", 105, 111, None,
        "detail", custom_headers=_pairs((2, "Eigandi"), (4, "Kennitala eiganda / land"), (5, "Hlutur"), (6, "Hlutfall")),
    ),
    _SectionSpec(
        "Hluthafaskrá", "C — Hlutafélög ekki áreiðanleikakönnuð", 118, 185, 117,
        "company", "company", 2, 3,
        custom_headers=_pairs(
            (1, "Nr"), (2, "Kennitala"), (3, "Nafn"), (4, "Hlutafé"),
            (5, "Hluthafi 1"), (6, "Hlutur 1"), (7, "Hluthafi 2"), (8, "Hlutur 2"),
            (9, "Hluthafi 3"), (10, "Hlutur 3"), (11, "Hluthafi 4"), (12, "Hlutur 4"),
            (14, "Arður 23"), (15, "Fjármtsk"), (16, "Framkv.stj"), (17, "Sími"),
            (18, "Netfang"), (19, "Netfang 2"),
        ),
        owner_pairs=_pairs((5, 6), (7, 8), (9, 10), (11, 12)),
        warning="Fyrirsögn í frumgögnum er skemmd og dálkar í nokkrum viðbættum línum eru hliðraðir.",
    ),
    _SectionSpec(
        "Hluthafaskrá", "D — SLF félög — ekki áreiðanleikakönnuð", 193, 198, 192,
        "slf", "slf", 2, 3, owner_pairs=_pairs((5, None), (6, None)),
    ),
    _SectionSpec(
        "Hluthafaskrá", "E — Félög sem Anna Bragina færir og ber ábyrgð á", 205, 208, 204,
        "company", "company", 2, 3, owner_pairs=_pairs((5, None), (6, None)),
    ),
    _SectionSpec(
        "Hluthafaskrá", "E — Félög sem Fribba, Íris og Alli færa og bera ábyrgð á", 213, 217, 212,
        "company", "company", 2, 3, owner_pairs=_pairs((5, None), (6, None)),
    ),
    _SectionSpec(
        "Hluthafaskrá", "F — Einstaklingar með rekstur", 222, 230, 221,
        "individual", "individual", 2, 3,
    ),
    _SectionSpec(
        "Hluthafaskrá", "G — Einstaklingar með rekstur — ekki áreiðanleikakannað", 237, 253, 236,
        "individual", "individual", 2, 3,
    ),

    _SectionSpec(
        "Áreiðanleikakönnun allir", "A — Hlutafélög í áreiðanleikakönnun", 4, 157, 3,
        "company", "company", 2, 3, owner_pairs=_pairs((6, 7), (8, 9), (10, 11), (12, 13)),
    ),
    _SectionSpec(
        "Áreiðanleikakönnun allir", "Sundurliðun Búferla", 160, 160, None,
        "company", "company", 2, 3,
        custom_headers=_pairs((1, "Sundurliðun"), (2, "Kennitala"), (3, "Nafn"), (6, "Hlutafé")),
    ),
    _SectionSpec(
        "Áreiðanleikakönnun allir", "Nánari sundurliðun Búferla", 162, 168, None,
        "detail", custom_headers=_pairs((1, "Eigandi"), (3, "Kennitala eiganda / land"), (6, "Hlutur"), (7, "Hlutfall")),
    ),
    _SectionSpec(
        "Áreiðanleikakönnun allir", "Unnið fyrir þriðja aðila — ekki áreiðanleikakannað", 176, 184, None,
        "company", "company", 2, 3,
        custom_headers=_pairs(
            (1, "Nr"), (2, "Kennitala"), (3, "Nafn"), (4, "Veflykill"),
            (6, "Hlutafé"), (7, "Eigandi 1"), (8, "Eigandi 2"), (9, "Netfang"),
            (10, "Sími"), (11, "Umsjón"),
        ), owner_pairs=_pairs((7, None), (8, None)),
    ),
    _SectionSpec(
        "Áreiðanleikakönnun allir", "SLF félög", 190, 193, 189,
        "slf", "slf", 2, 3, owner_pairs=_pairs((6, None), (7, None)),
    ),
    _SectionSpec(
        "Áreiðanleikakönnun allir", "F — Einstaklingar með rekstur", 199, 223, 198,
        "individual", "individual", 2, 3,
    ),
)


SECTION_LABEL_RANGES = {
    "Hlutafjármiðar 2025": (
        (1, 189, "Hlutafjármiðar 2025"),
        (190, 199, "Sundurliðun Búferla og ómerktir útreikningar"),
    ),
    "Hluthafaskrá": (
        (1, 99, "A — Hlutafélög í áreiðanleikakönnun"),
        (100, 114, "B — Nánari sundurliðun"),
        (115, 188, "C — Hlutafélög ekki áreiðanleikakönnuð"),
        (189, 201, "D — SLF félög — ekki áreiðanleikakönnuð"),
        (202, 209, "E — Félög sem Anna Bragina færir og ber ábyrgð á"),
        (210, 218, "E — Félög sem Fribba, Íris og Alli færa og bera ábyrgð á"),
        (219, 233, "F — Einstaklingar með rekstur"),
        (234, 253, "G — Einstaklingar með rekstur — ekki áreiðanleikakannað"),
    ),
    "Áreiðanleikakönnun allir": (
        (1, 173, "A — Hlutafélög í áreiðanleikakönnun"),
        (174, 186, "Unnið fyrir þriðja aðila — ekki áreiðanleikakannað"),
        (187, 195, "SLF félög"),
        (196, 226, "F — Einstaklingar með rekstur"),
        (227, 231, "Samantekt"),
    ),
}

TITLE_ROWS = {
    "Hlutafjármiðar 2025": {1},
    "Hluthafaskrá": {1, 100, 115, 189, 202, 210, 219, 234},
    "Áreiðanleikakönnun allir": {1, 174, 187, 196},
}


_ICELANDIC_TRANSLATION = str.maketrans({"ð": "d", "þ": "th", "æ": "ae", "ö": "o"})
_EMAIL_RE = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.IGNORECASE)


def normalize_text(value) -> str:
    text = str(value or "").strip().casefold().translate(_ICELANDIC_TRANSLATION)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(re.findall(r"[a-z0-9]+", text))


def normalize_kennitala(value) -> str:
    digits = "".join(re.findall(r"\d", str(value or "")))
    return digits if len(digits) == 10 else ""


def _section_spec(sheet: str, row: int) -> _SectionSpec | None:
    for spec in SECTION_SPECS:
        if spec.sheet == sheet and spec.start_row <= row <= spec.end_row:
            return spec
    return None


def _section_label(sheet: str, row: int) -> str:
    spec = _section_spec(sheet, row)
    if spec is not None:
        return spec.section
    for start, end, label in SECTION_LABEL_RANGES[sheet]:
        if start <= row <= end:
            return label
    return "Óflokkuð frumgögn"


def _header_map(workbook, spec: _SectionSpec | None) -> dict[int, str]:
    headers = {}
    if spec is None:
        return headers
    if spec.header_row:
        worksheet = workbook[spec.sheet]
        max_column = SHEET_BOUNDS[spec.sheet][1]
        for column in range(1, max_column + 1):
            value = worksheet.cell(spec.header_row, column).value
            if value not in (None, ""):
                headers[column] = str(value).strip()
    headers.update(dict(spec.custom_headers))
    return headers


def _is_sensitive(header: str, row: int, spec: _SectionSpec | None) -> bool:
    if spec is None or row == spec.header_row:
        return False
    label = normalize_text(header)
    return any(token in label for token in ("veflykill", "lykilord", "password", "credential", "access key"))


def _cell_value(cell):
    value = cell.value
    if hasattr(value, "text") and value.__class__.__name__ in {"ArrayFormula", "DataTableFormula"}:
        return value.text
    return value


def _read_source_rows(path: str | Path) -> tuple[SourceRow, ...]:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)

    formula_book = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    missing = [name for name in TARGET_SHEETS if name not in formula_book.sheetnames]
    if missing:
        formula_book.close()
        raise RegistryWorkbookError(
            "Í völdu vinnubókinni vantar nauðsynleg blöð hluthafaskrárinnar: " + ", ".join(missing)
        )
    cached_book = load_workbook(path, data_only=True, read_only=False, keep_links=True)

    rows = []
    try:
        for sheet in TARGET_SHEETS:
            formula_sheet = formula_book[sheet]
            cached_sheet = cached_book[sheet]
            max_row, max_column = SHEET_BOUNDS[sheet]
            by_row = defaultdict(list)

            # _cells iterates only materialized XML cells.  This is deliberate:
            # iter_rows/max_row would walk phantom XFD/bottom-right dimensions in
            # other sheets of this workbook.
            for (row_index, column_index), cell in list(formula_sheet._cells.items()):
                if row_index > max_row or column_index > max_column:
                    continue
                original = _cell_value(cell)
                if original is None:
                    continue
                spec = _section_spec(sheet, row_index)
                headers = _header_map(formula_book, spec)
                header = headers.get(column_index) or f"Án fyrirsagnar ({get_column_letter(column_index)})"
                is_formula = cell.data_type == "f" or (
                    isinstance(original, str) and original.startswith("=")
                )
                formula = str(original) if is_formula else None
                cached = _cell_value(cached_sheet.cell(row_index, column_index)) if is_formula else None
                value = cached if is_formula and cached is not None else original
                hyperlink = None
                if cell.hyperlink is not None:
                    hyperlink = cell.hyperlink.target or cell.hyperlink.location
                by_row[row_index].append(
                    SourceCell(
                        sheet=sheet,
                        section=_section_label(sheet, row_index),
                        excel_row=row_index,
                        column=get_column_letter(column_index),
                        column_index=column_index,
                        header=header,
                        original_value=original,
                        formula=formula,
                        cached_value=cached,
                        value=value,
                        sensitive=_is_sensitive(header, row_index, spec),
                        hyperlink=hyperlink,
                    )
                )

            header_rows = {spec.header_row for spec in SECTION_SPECS if spec.sheet == sheet and spec.header_row}
            for row_index in sorted(by_row):
                spec = _section_spec(sheet, row_index)
                if row_index in header_rows or row_index in TITLE_ROWS[sheet]:
                    record_type = "heading"
                else:
                    record_type = spec.record_type if spec else "raw"
                warning = spec.warning if spec else ""
                rows.append(
                    SourceRow(
                        sheet=sheet,
                        section=_section_label(sheet, row_index),
                        excel_row=row_index,
                        cells=tuple(sorted(by_row[row_index], key=lambda item: item.column_index)),
                        record_type=record_type,
                        layout_warning=warning,
                    )
                )
    finally:
        formula_book.close()
        cached_book.close()

    return tuple(rows)


def _row_cell(row: SourceRow, column_index: int) -> SourceCell | None:
    return next((cell for cell in row.cells if cell.column_index == column_index), None)


def _row_value(row: SourceRow, column_index: int):
    cell = _row_cell(row, column_index)
    return cell.value if cell is not None else None


def _clean_name(value) -> str:
    return str(value or "").replace("\xa0", " ").strip()


def _unique(values) -> tuple:
    output = []
    seen = set()
    for value in values:
        if value in (None, ""):
            continue
        marker = normalize_text(value) if isinstance(value, str) else value
        if marker in seen:
            continue
        seen.add(marker)
        output.append(value)
    return tuple(output)


@dataclass
class _Observation:
    row: SourceRow
    spec: _SectionSpec
    kennitala: str
    name: str


@dataclass
class _EntityBuilder:
    key: str
    observations: list[_Observation] = field(default_factory=list)
    detail_rows: list[SourceRow] = field(default_factory=list)
    related_parties: list[RelatedParty] = field(default_factory=list)
    reverse_ownerships: list[RelatedParty] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _collect_observations(rows: tuple[SourceRow, ...]) -> list[_Observation]:
    observations = []
    for row in rows:
        spec = _section_spec(row.sheet, row.excel_row)
        if spec is None or not spec.is_entity:
            continue
        name = _clean_name(_row_value(row, spec.name_column))
        if not name:
            continue
        identifier = normalize_kennitala(_row_value(row, spec.id_column)) if spec.id_column else ""
        entity_row = row
        if row.sheet == "Hlutafjármiðar 2025" and row.excel_row == 191:
            entity_row = replace(
                row,
                cells=tuple(cell for cell in row.cells if cell.column_index <= 6),
                layout_warning="Ótengdir reitir I–K eru áfram aðgengilegir í yfirliti frumgagna.",
            )
        observations.append(_Observation(row=entity_row, spec=spec, kennitala=identifier, name=name))
    return observations


def _canonical_keys(observations: list[_Observation]) -> dict[int, str]:
    names_to_ids = defaultdict(set)
    for observation in observations:
        if observation.kennitala:
            names_to_ids[normalize_text(observation.name)].add(observation.kennitala)

    keys = {}
    for observation in observations:
        if observation.kennitala:
            key = f"kt:{observation.kennitala}"
        else:
            matching_ids = names_to_ids.get(normalize_text(observation.name), set())
            if len(matching_ids) == 1:
                key = f"kt:{next(iter(matching_ids))}"
            elif not matching_ids:
                key = f"name:{normalize_text(observation.name)}"
            else:
                key = f"row:{observation.row.sheet}:{observation.row.excel_row}"
        keys[id(observation)] = key
    return keys


def _value_kind(value) -> str:
    text = str(value or "").strip()
    if not text:
        return "empty"
    if normalize_kennitala(text):
        return "kennitala"
    if _EMAIL_RE.search(text):
        return "email"
    try:
        float(text.replace(".", "").replace(",", "."))
        return "number"
    except (TypeError, ValueError):
        return "text"


def _related_party_from_pair(row: SourceRow, owner_column: int, share_column: int | None, ordinal: int) -> RelatedParty | None:
    owner_cell = _row_cell(row, owner_column)
    share_cell = _row_cell(row, share_column) if share_column else None
    owner_value = owner_cell.value if owner_cell else None
    share_value = share_cell.value if share_cell else None
    if owner_value in (None, "") and share_value in (None, ""):
        return None

    warning = ""
    owner_kind = _value_kind(owner_value)
    share_kind = _value_kind(share_value)
    if owner_kind == "empty" and share_kind == "kennitala":
        owner_value, share_value = share_value, None
        owner_kind = "kennitala"
        warning = "Kennitala eiganda var skráð í dálk fyrir hlut."
    elif owner_kind == "empty" and share_kind == "text":
        owner_value, share_value = share_value, None
        owner_kind = "text"
        warning = "Nafn eiganda var skráð í dálk fyrir hlut."
    elif owner_kind == "empty":
        return None
    elif owner_kind == "number" and share_kind in {"kennitala", "text"}:
        owner_value, share_value = share_value, owner_value
        owner_kind = share_kind
        warning = "Dálkar fyrir eiganda og hlut virðast hliðraðir; gildi voru pöruð eftir tegund."
    elif owner_kind == "number":
        return None

    identifier = normalize_kennitala(owner_value)
    name = "" if identifier else _clean_name(owner_value)
    return RelatedParty(
        role=f"Hluthafi {ordinal}",
        name=name,
        kennitala=identifier,
        share=share_value,
        source_sheet=row.sheet,
        excel_row=row.excel_row,
        warning=warning,
    )


def _detail_related_party(row: SourceRow) -> RelatedParty | None:
    if row.sheet == "Hlutafjármiðar 2025":
        name_column, id_column, share_column, percent_column = 1, 3, 5, 6
    elif row.sheet == "Hluthafaskrá":
        name_column, id_column, share_column, percent_column = 2, 4, 5, 6
    else:
        name_column, id_column, share_column, percent_column = 1, 3, 6, 7
    name = _clean_name(_row_value(row, name_column))
    raw_identifier = _row_value(row, id_column)
    identifier = normalize_kennitala(raw_identifier)
    if not name and not raw_identifier:
        return None
    warning = "" if identifier or not raw_identifier else f"Auðkenni eiganda er land eða athugasemd en ekki kennitala: {raw_identifier}."
    share = _row_value(row, share_column)
    if isinstance(share, str) and share.startswith("#"):
        warning = (warning + " " if warning else "") + "Vistað niðurstöðugildi formúlunnar fyrir hlut inniheldur Excel-villu."
    return RelatedParty(
        role="Eigandi í sundurliðun",
        name=name,
        kennitala=identifier,
        share=share,
        percentage=_row_value(row, percent_column),
        source_sheet=row.sheet,
        excel_row=row.excel_row,
        warning=warning,
    )


def _extract_contacts(rows: list[SourceRow]):
    emails = []
    phones = []
    for row in rows:
        for cell in row.cells:
            if cell.sensitive:
                continue
            text = str(cell.value or "").strip().replace("\xa0", " ")
            label = normalize_text(cell.header)
            email_match = _EMAIL_RE.search(text)
            if email_match:
                emails.append(email_match.group(0).rstrip(">"))
            elif any(token in label for token in ("simi", "phone")):
                digits = "".join(re.findall(r"\d", text))
                if 6 <= len(digits) <= 10:
                    phones.append(text)
    return _unique(emails), _unique(phones)


def _normalized_observation_value(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip().replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return normalize_text(value)


def _data_quality_warnings(builder: _EntityBuilder) -> list[str]:
    warnings = list(builder.warnings)
    observations = builder.observations
    if not any(observation.kennitala for observation in observations):
        warnings.append("Engin gild kennitala er skráð; færslan er áfram tengd við upprunalínu sína.")
    per_sheet = defaultdict(list)
    for observation in observations:
        per_sheet[observation.row.sheet].append(observation)
        if observation.row.layout_warning:
            warnings.append(observation.row.layout_warning)
    for sheet, values in per_sheet.items():
        if len(values) > 1:
            warnings.append(f"{len(values)} upprunalínur hafa sama auðkenni í {sheet}; engin þeirra var yfirskrifuð.")

    names = _unique(observation.name for observation in observations)
    if len(names) > 1:
        warnings.append("Nafn sama aðila er ritað á fleiri en einn hátt í upprunablöðunum.")

    capital_values = defaultdict(list)
    credential_values = []
    for observation in observations:
        for cell in observation.row.cells:
            label = normalize_text(cell.header)
            if "hlutafe" in label or "stofnfe" in label:
                if cell.value not in (None, ""):
                    capital_values[_normalized_observation_value(cell.value)].append(cell)
            if cell.sensitive and cell.value not in (None, ""):
                credential_values.append(str(cell.value).strip().casefold())
    if len(capital_values) > 1:
        warnings.append("Gildi hlutafjár stangast á milli upprunafærslna.")
    if len(set(credential_values)) > 1:
        warnings.append("Veflyklar stangast á milli upprunafærslna; hvert gildi er áfram merkt uppruna sínum.")
    return list(_unique(warnings))


class ShareholderRegistry:
    """Immutable, source-preserving registry with ranked public-field search."""

    def __init__(self, path: str | Path, source_rows: tuple[SourceRow, ...], entities: tuple[RegistryEntity, ...]):
        self.path = str(Path(path).resolve())
        self.source_rows = source_rows
        self.entities = tuple(sorted(entities, key=lambda item: normalize_text(item.name)))
        self._by_key = {entity.key: entity for entity in self.entities}
        self.sheet_names = TARGET_SHEETS
        sheet_counts = {
            sheet: sum(len(row.cells) for row in source_rows if row.sheet == sheet)
            for sheet in TARGET_SHEETS
        }
        all_cells = tuple(cell for row in source_rows for cell in row.cells)
        self.stats = RegistryStats(
            entity_count=len(self.entities),
            company_count=sum(entity.kind in {"company", "slf"} for entity in self.entities),
            individual_count=sum(entity.kind == "individual" for entity in self.entities),
            source_row_count=len(source_rows),
            source_cell_count=len(all_cells),
            sensitive_cell_count=sum(cell.sensitive for cell in all_cells),
            formula_cell_count=sum(cell.formula is not None for cell in all_cells),
            sheet_counts=sheet_counts,
        )

    @property
    def raw_cells(self) -> tuple[SourceCell, ...]:
        return tuple(cell for row in self.source_rows for cell in row.cells)

    def get(self, key: str) -> RegistryEntity | None:
        if key in self._by_key:
            return self._by_key[key]
        digits = normalize_kennitala(key)
        return self._by_key.get(f"kt:{digits}") if digits else None

    def iter_source_rows(self, sheet=None, section=None, entity_key=None):
        if entity_key:
            entity = self.get(entity_key)
            rows = () if entity is None else (*entity.source_rows, *entity.detail_rows)
        else:
            rows = self.source_rows
        for row in rows:
            if sheet and row.sheet != sheet:
                continue
            if section and row.section != section:
                continue
            yield row

    def search(self, query: str, limit: int = 50) -> tuple[SearchResult, ...]:
        query = str(query or "").strip()
        normalized_query = normalize_text(query)
        query_digits = "".join(re.findall(r"\d", query))
        phone_query = query_digits if len(query_digits) >= 4 and not re.search(r"[A-Za-zÀ-ÿ@]", query) else ""
        if not normalized_query:
            return tuple(
                SearchResult(entity, 0, "", "")
                for entity in self.entities[:max(0, int(limit))]
            )

        tokens = normalized_query.split()
        results = []
        for entity in self.entities:
            name = normalize_text(entity.name)
            aliases = [normalize_text(alias) for alias in entity.aliases]
            score = 0
            matched_on = ""
            matched_value = ""
            if query_digits and len(query_digits) >= 4 and query_digits == entity.kennitala:
                score, matched_on, matched_value = 1000, "kennitala", entity.kennitala
            elif normalized_query == name:
                score, matched_on, matched_value = 950, "name", entity.name
            elif name.startswith(normalized_query):
                score, matched_on, matched_value = 850, "name prefix", entity.name
            elif tokens and all(token in name for token in tokens):
                score, matched_on, matched_value = 780, "name", entity.name
            elif normalized_query in name:
                score, matched_on, matched_value = 720, "name", entity.name
            else:
                for alias, normalized_alias in zip(entity.aliases, aliases):
                    if normalized_query == normalized_alias or normalized_query in normalized_alias:
                        score, matched_on, matched_value = 680, "alias", alias
                        break
            if not score:
                for email in entity.emails:
                    if normalized_query in normalize_text(email):
                        score, matched_on, matched_value = 520, "email", email
                        break
            if not score:
                for phone in entity.phones:
                    phone_digits = "".join(re.findall(r"\d", phone))
                    if phone_query and phone_query in phone_digits:
                        score, matched_on, matched_value = 500, "phone", phone
                        break
            if not score:
                for party in entity.related_parties:
                    party_name = normalize_text(party.name)
                    if (
                        (normalized_query and normalized_query in party_name)
                        or (len(query_digits) == 10 and query_digits == normalize_kennitala(party.kennitala))
                    ):
                        score = 460
                        matched_on = "shareholder"
                        matched_value = party.name or party.kennitala
                        break
            if score:
                results.append(SearchResult(entity, score, matched_on, matched_value))
        results.sort(key=lambda result: (-result.score, normalize_text(result.entity.name)))
        return tuple(results[:max(0, int(limit))])


def _build_registry(path: str | Path, rows: tuple[SourceRow, ...]) -> ShareholderRegistry:
    observations = _collect_observations(rows)
    keys = _canonical_keys(observations)
    builders: dict[str, _EntityBuilder] = {}
    for observation in observations:
        key = keys[id(observation)]
        builder = builders.setdefault(key, _EntityBuilder(key))
        builder.observations.append(observation)

    buferlar = builders.get("kt:5304221190")
    if buferlar is not None:
        for row in rows:
            # All explicit detail sections in this workbook describe the
            # Búferlar ownership breakdown.  Keying this attachment by the
            # structural record type also keeps it robust across console and
            # workbook Unicode normalisation differences.
            if row.record_type == "detail":
                buferlar.detail_rows.append(row)
                party = _detail_related_party(row)
                if party is not None:
                    buferlar.related_parties.append(party)

    for builder in builders.values():
        for observation in builder.observations:
            for ordinal, (owner_column, share_column) in enumerate(observation.spec.owner_pairs, start=1):
                party = _related_party_from_pair(observation.row, owner_column, share_column, ordinal)
                if party is not None:
                    builder.related_parties.append(party)

    # Reverse links are intentionally limited to exact normalized kennitölur.
    for target_builder in builders.values():
        target_ids = [item.kennitala for item in target_builder.observations if item.kennitala]
        target_id = target_ids[0] if target_ids else ""
        target_name = target_builder.observations[0].name if target_builder.observations else ""
        for relationship in target_builder.related_parties:
            owner_key = f"kt:{normalize_kennitala(relationship.kennitala)}"
            owner_builder = builders.get(owner_key)
            if owner_builder is None:
                continue
            owner_builder.reverse_ownerships.append(
                RelatedParty(
                    role="Á hlut í",
                    name=target_name,
                    kennitala=target_id,
                    share=relationship.share,
                    percentage=relationship.percentage,
                    source_sheet=relationship.source_sheet,
                    excel_row=relationship.excel_row,
                    warning=relationship.warning,
                )
            )

    entities = []
    for key, builder in builders.items():
        observations_sorted = sorted(
            builder.observations,
            key=lambda item: (SHEET_PRIORITY.get(item.row.sheet, 99), item.row.excel_row),
        )
        identifiers = _unique(item.kennitala for item in observations_sorted if item.kennitala)
        names = _unique(item.name for item in observations_sorted)
        kinds = [item.spec.entity_kind for item in observations_sorted]
        if kinds and all(kind == "individual" for kind in kinds):
            kind = "individual"
        elif kinds and all(kind == "slf" for kind in kinds):
            kind = "slf"
        else:
            kind = "company"
        source_rows = [item.row for item in observations_sorted]
        emails, phones = _extract_contacts(source_rows)
        statuses = _unique(item.spec.section for item in observations_sorted)
        warnings = _data_quality_warnings(builder)
        entities.append(
            RegistryEntity(
                key=key,
                kind=kind,
                kennitala=identifiers[0] if identifiers else "",
                name=names[0] if names else "Ónefnd færsla",
                aliases=tuple(names[1:]),
                emails=emails,
                phones=phones,
                statuses=statuses,
                source_rows=tuple(source_rows),
                related_parties=tuple(builder.related_parties),
                reverse_ownerships=tuple(builder.reverse_ownerships),
                detail_rows=tuple(builder.detail_rows),
                warnings=tuple(warnings),
            )
        )

    return ShareholderRegistry(path, rows, tuple(entities))


def load_registry(path: str | Path) -> ShareholderRegistry:
    """Load and link the three intended registry sheets from *path*."""
    rows = _read_source_rows(path)
    return _build_registry(path, rows)


load_registry_data = load_registry


class LazyShareholderRegistry:
    """Thread-safe one-time loader used by UI integrations and tests."""

    def __init__(self, path: str | Path):
        self.path = str(path)
        self.registry: ShareholderRegistry | None = None
        self._lock = threading.Lock()

    @property
    def is_loaded(self) -> bool:
        return self.registry is not None

    def load(self) -> ShareholderRegistry:
        if self.registry is None:
            with self._lock:
                if self.registry is None:
                    self.registry = load_registry(self.path)
        return self.registry

    def search(self, query: str, limit: int = 50) -> tuple[SearchResult, ...]:
        return self.load().search(query, limit=limit)
