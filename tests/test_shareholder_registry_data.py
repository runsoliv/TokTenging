from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from tok.shareholder_registry_data import (
    LazyShareholderRegistry,
    RegistryWorkbookError,
    TARGET_SHEETS,
    load_registry,
)


HFM_HEADERS = [
    "Nr", "Kennitala", "Nafn", "Veflykill", "Hlutafé", "Hluthafi 1", "Hlutur",
    "Hluthafi 2", "Hlutur", "Hluthafi 3", "Hlutur", "Hluthafi 4", "Hlutur",
    "Arður 25", "Fjármtsk",
]
HK_HEADERS = [
    "Nr", "Kennitala", "Nafn", "Hlutafé", "Hluthafi 1", "Hlutur", "Hluthafi 2",
    "Hlutur", "Hluthafi 3", "Hlutur", "Hluthafi 4", "Hlutur", "Arður 23",
    "Fjármtsk", "Framkv.stj", "Sími", "Netfang",
]
AR_HEADERS = [
    "Nr", "Kennitala", "Nafn", "Veflykill", "Hlutafé", "Hluthafi 1", "Hlutur",
    "Hluthafi 2", "Hlutur", "Hluthafi 3", "Hlutur", "Hluthafi 4", "Hlutur",
    "Arður 23", "Fjármtsk", "Framkv.stj", "Sími", "Netfang", "netfang2",
]


def build_fixture(path):
    workbook = Workbook()
    hfm = workbook.active
    hfm.title = TARGET_SHEETS[0]
    hk = workbook.create_sheet(TARGET_SHEETS[1])
    ar = workbook.create_sheet(TARGET_SHEETS[2])
    other = workbook.create_sheet("Unrelated")

    for sheet, title, headers in (
        (hfm, "Hlutafjármiðar 2025", HFM_HEADERS),
        (hk, "A — Companies", HK_HEADERS),
        (ar, "A — Companies", AR_HEADERS),
    ):
        sheet["A1"] = title
        for column, header in enumerate(headers, 1):
            sheet.cell(3, column, header)

    hfm.append([])
    for column, value in enumerate(
        (1, "5709101910", "Þakmálun ehf.", "SECRET-ONLY", 500000, "010101-1234", 500000),
        1,
    ):
        hfm.cell(4, column, value)
    hfm["P4"] = "=E4/2"
    hfm["C5"] = "Orphan ehf"

    for column, value in enumerate(
        (1, "570910-1910", "Þakmálun ehf", 500000, "0101011234", 500000),
        1,
    ):
        hk.cell(4, column, value)
    hk["Q4"] = "office@example.is"
    hk["Q4"].hyperlink = "mailto:registry@example.is"

    for column, value in enumerate(
        (1, "5709101910", "Þakmálun ehf.", "SECOND-SECRET", 500000, "010101-1234", 500000),
        1,
    ):
        ar.cell(4, column, value)
    ar["R4"] = "office@example.is"
    ar["C5"] = "Orphan ehf"

    # These cells prove that stale/far-edge dimensions and unrelated sheets are
    # not allowed to inflate the registry import.
    hfm["XFD1048576"] = 999
    other["XFD1048576"] = "not registry data"
    workbook.save(path)


class GeneratedWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "registry.xlsx"
        build_fixture(self.path)
        self.registry = load_registry(self.path)

    def tearDown(self):
        self.temp.cleanup()

    def test_preserves_bounded_source_cells_formulas_and_hyperlinks(self):
        self.assertEqual(self.registry.stats.source_cell_count, 79)
        self.assertEqual(self.registry.stats.source_row_count, 11)
        self.assertEqual(self.registry.stats.formula_cell_count, 1)
        self.assertEqual(self.registry.stats.sensitive_cell_count, 2)
        self.assertFalse(any(cell.coordinate == "XFD1048576" for cell in self.registry.raw_cells))

        formula = next(cell for cell in self.registry.raw_cells if cell.coordinate == "P4" and cell.sheet == TARGET_SHEETS[0])
        self.assertEqual(formula.formula, "=E4/2")
        self.assertEqual(formula.original_value, "=E4/2")

        linked = next(cell for cell in self.registry.raw_cells if cell.sheet == TARGET_SHEETS[1] and cell.coordinate == "Q4")
        self.assertEqual(linked.hyperlink, "mailto:registry@example.is")

    def test_merges_by_id_and_unambiguous_name_fallback(self):
        self.assertEqual(self.registry.stats.entity_count, 2)
        company = self.registry.get("570910-1910")
        self.assertIsNotNone(company)
        self.assertEqual(len(company.source_rows), 3)
        self.assertEqual(company.kennitala, "5709101910")

        orphan = self.registry.search("Orphan ehf", limit=5)[0].entity
        self.assertTrue(orphan.key.startswith("name:"))
        self.assertEqual(len(orphan.source_rows), 2)
        self.assertEqual(orphan.kennitala, "")

    def test_search_is_icelandic_aware_and_excludes_credentials(self):
        self.assertEqual(self.registry.search("Thakmalun")[0].entity.name, "Þakmálun ehf.")
        self.assertEqual(self.registry.search("570910-1910")[0].matched_on, "kennitala")
        self.assertEqual(self.registry.search("office@example.is")[0].matched_on, "email")
        self.assertEqual(self.registry.search("010101-1234")[0].matched_on, "shareholder")
        self.assertEqual(self.registry.search("SECRET-ONLY"), ())
        self.assertEqual(self.registry.search("SECOND-SECRET"), ())

        secret = next(cell for cell in self.registry.raw_cells if cell.sensitive)
        self.assertEqual(secret.display_value(), "••••••••")
        self.assertNotEqual(secret.display_value(reveal_sensitive=True), "••••••••")

    def test_lazy_loader_and_entity_row_filter(self):
        lazy = LazyShareholderRegistry(self.path)
        self.assertFalse(lazy.is_loaded)
        result = lazy.search("Thakmalun")[0]
        self.assertTrue(lazy.is_loaded)
        rows = list(lazy.registry.iter_source_rows(entity_key=result.entity.key))
        self.assertEqual(len(rows), 3)

    def test_missing_target_sheet_is_rejected_cleanly(self):
        broken = Path(self.temp.name) / "broken.xlsx"
        workbook = Workbook()
        workbook.save(broken)
        with self.assertRaises(RegistryWorkbookError):
            load_registry(broken)


class RealWorkbookIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.path = (
            Path(__file__).resolve().parents[1]
            / "tok"
            / "trainingData"
            / "data folder"
            / "data codex.xlsx"
        )
        if not cls.path.is_file():
            raise unittest.SkipTest("Private registry workbook is not present")
        cls.registry = load_registry(cls.path)

    def test_exact_source_completeness_and_privacy_counts(self):
        stats = self.registry.stats
        self.assertEqual(stats.entity_count, 237)
        self.assertEqual(stats.source_row_count, 635)
        self.assertEqual(stats.source_cell_count, 4503)
        self.assertEqual(stats.formula_cell_count, 144)
        self.assertEqual(stats.sensitive_cell_count, 294)
        self.assertEqual(
            stats.sheet_counts,
            {TARGET_SHEETS[0]: 1316, TARGET_SHEETS[1]: 1564, TARGET_SHEETS[2]: 1623},
        )
        self.assertEqual(sum(bool(cell.hyperlink) for cell in self.registry.raw_cells), 295)

    def test_representative_company_owner_and_missing_id_searches(self):
        self.assertEqual(self.registry.search("300Brains")[0].entity.kennitala, "5802191520")
        self.assertEqual(self.registry.search("Thakmalun")[0].entity.name, "Þakmálun ehf.")
        self.assertTrue(self.registry.search("Glide ride")[0].entity.key.startswith("name:"))
        owner_matches = {result.entity.name for result in self.registry.search("490922-0830", limit=20)}
        self.assertIn("300Brains ehf", owner_matches)
        self.assertIn("Búferlar ehf", owner_matches)

    def test_formula_cache_and_secret_search_invariants(self):
        formulas = [cell for cell in self.registry.raw_cells if cell.formula]
        self.assertEqual(len(formulas), 144)
        self.assertTrue(all(cell.cached_value is not None for cell in formulas))
        secrets = [cell for cell in self.registry.raw_cells if cell.sensitive]
        self.assertEqual(len(secrets), 294)
        for cell in secrets:
            self.assertEqual(cell.display_value(), "••••••••")
        for cell in secrets[:20]:
            self.assertEqual(self.registry.search(str(cell.value), limit=500), ())


if __name__ == "__main__":
    unittest.main()
