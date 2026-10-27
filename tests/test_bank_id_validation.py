import unittest
from pathlib import Path
import tempfile

from openpyxl import Workbook
import pandas as pd

from tok.bank_detection import BANK_CONFIGS, detect_bank_type
from tok.bank_formatter import _arion_amount_source, _prepare_bank_excel_output
from tok.bank_utils import _format_kennitala


class KennitalaValidationTests(unittest.TestCase):
    def test_accepts_only_the_two_supported_kennitala_shapes(self):
        accepted = {
            "1234567890": "1234567890",
            "123456-7890": "1234567890",
            1234567890: "1234567890",
            1234567890.0: "1234567890",
        }
        for value, expected in accepted.items():
            with self.subTest(value=value):
                self.assertEqual(_format_kennitala(value), expected)

    def test_rejects_references_and_non_ten_digit_values(self):
        rejected = (
            None,
            "",
            "123456789",
            "12345678901",
            "123 456 7890",
            "12345-67890",
            "REF-123456",
            123456789,
            1234567890.5,
            True,
        )
        for value in rejected:
            with self.subTest(value=value):
                self.assertEqual(_format_kennitala(value), "")

    def test_final_output_guard_blanks_invalid_ids(self):
        source = pd.DataFrame({
            "DATE": ["01.01.2026"] * 5,
            "TEXT": ["Færsla"] * 5,
            "ID": ["1234567890", "123456-7890", "ABC-123", "123456789", None],
            "AMOUNT": [1] * 5,
        })

        output = _prepare_bank_excel_output(source)

        self.assertEqual(
            output["ID"].tolist(),
            ["1234567890", "1234567890", "", "", ""],
        )

    def test_islandsbanki_account_exports_use_kennitala_columns(self):
        standard = BANK_CONFIGS["islandsbanki_reikningsyfirlit"]
        metadata = BANK_CONFIGS["islandsbanki_reikningsyfirlit_metadata"]

        self.assertIn("Kennitala mótaðila", standard["columns"])
        self.assertEqual(standard["rename"]["Kennitala mótaðila"], "ID")
        self.assertNotIn("Tilvísun", standard["rename"])
        self.assertIn("Kennitala", metadata["columns"])
        self.assertEqual(metadata["rename"]["Kennitala"], "ID")
        self.assertNotIn("Tilvísun", metadata["rename"])

    def test_detection_preserves_a_leading_zero_in_kennitala(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "islandsbanki.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "Dagsetning",
                "Mótaðili",
                "Kennitala mótaðila",
                "Tilvísun",
                "Texti",
                "Upphæð",
                "Staða",
            ])
            sheet.append([
                "01.01.2026",
                "Prófun ehf.",
                "0123456789",
                "REF-123",
                "Prófun",
                100,
                "Bókað",
            ])
            workbook.save(path)

            bank_type, data = detect_bank_type(path)

        self.assertEqual(bank_type, "islandsbanki_reikningsyfirlit")
        self.assertEqual(data.loc[0, "Kennitala mótaðila"], "0123456789")

    def test_detection_handles_new_islandsbanki_metadata_statement(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "reikningsyfirlit_2026-07-28.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            metadata_rows = [
                ["Eigandi", "Prófun ehf."],
                ["Kennitala", "4111220850"],
                ["Reikningur", "Almennur reikningur"],
                ["Reikningsnúmer", "0000-00-000000"],
                ["Staða", 100000],
                [],
                ["Dagsetning frá", "01.05.2026"],
                ["Dagsetning til", "30.06.2026"],
                ["Yfirlit sótt", "2026-07-28 11:09:18"],
                [],
                [],
            ]
            for row in metadata_rows:
                sheet.append(row)
            sheet.append([
                "Dagsetning",
                "Mótaðili",
                "Kennitala mótaðila",
                "Tilvísun",
                "Texti",
                "Upphæð",
                "Staða",
            ])
            sheet.append([
                "30.06.2026",
                "Prófun ehf.",
                "0123456789",
                "0000109",
                "Millifært",
                -150000,
                250000,
            ])
            workbook.save(path)

            bank_type, data = detect_bank_type(path)

        self.assertEqual(bank_type, "islandsbanki_reikningsyfirlit")
        self.assertEqual(len(data), 1)
        self.assertEqual(
            list(data.columns),
            [
                "Dagsetning",
                "Mótaðili",
                "Kennitala mótaðila",
                "Tilvísun",
                "Texti",
                "Upphæð",
                "Staða",
            ],
        )
        self.assertEqual(data.loc[0, "Kennitala mótaðila"], "0123456789")
        self.assertEqual(data.loc[0, "Upphæð"], -150000)

    def test_detection_handles_islandsbanki_statement_with_textalykill(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "islandsbanki_textalykill.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in (
                ["Reikningseigandi", "Linh þrif ehf."],
                ["Kennitala", "5910220400"],
                ["Reikningur", "Almennur reikningur"],
                ["Reikningsnúmer", "0515-26-010049"],
                ["Staða", 311592],
                [],
                ["Dagsetning frá", "01.05.2026"],
                ["Dagsetning til", "30.06.2026"],
                ["Yfirlit sótt", "2026-08-04 16:59:13"],
                [],
                [],
            ):
                sheet.append(row)
            sheet.append([
                "Dagsetning",
                "Mótaðili",
                "Kennitala mótaðila",
                "Tilvísun",
                "Textalykill",
                "Upphæð",
                "Staða",
            ])
            sheet.append([
                "30.06.2026",
                "Lyfja Lagmula",
                "0123456789",
                "0981060",
                "Úttekt með debetkorti",
                -15694,
                180844,
            ])
            workbook.save(path)

            bank_type, data = detect_bank_type(path)

        self.assertEqual(bank_type, "islandsbanki_reikningsyfirlit")
        self.assertEqual(len(data), 1)
        self.assertIn("Textalykill", data.columns)
        self.assertEqual(data.loc[0, "Kennitala mótaðila"], "0123456789")

    def test_islandsbanki_special_without_kennitala_stays_special(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "islandsbanki_special.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Dagsetning", "Mótaðili", "Upphæð"])
            sheet.append(["01.01.2026", "Sérsniðin færsla", -1000])
            workbook.save(path)

            bank_type, _ = detect_bank_type(path)

        self.assertEqual(bank_type, "islandsbanki_special")

    def test_detection_handles_new_arion_account_transactions_export(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "AccountTransactions.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in (
                ["Reikningseigandi", None, None, "Nánari upplýsingar"],
                ["Eigandi", "Prófun ehf."],
                ["Kennitala", "411122-0850"],
                [], ["Stofnupplýsingar"],
                ["Reikningsnúmer", "0308-26-004408"],
                ["Tegund reiknings", "Fyrirtækjareikningur"],
                ["IBAN númer", "IS26 0308 2600 4408 4408 1704 10"],
                ["SWIFT", "ESJAISRE"],
                [], ["Tímabil"],
                ["Tímabil frá", "01.05.2026"],
                ["Tímabil til", "30.06.2026"],
                [], [],
            ):
                sheet.append(row)
            sheet.append([
                "Dagsetning", "Upphæð", "Staða", "Mynt", "Skýring",
                "Seðilnúmer", "Tilvísun", "Texti", "Færslulykill",
                "Vaxtadagur", "Textalykill", "Greiðslubanki", "Bunkanúmer",
                "Vaxtadagsetning", "Nafn viðtakanda eða greiðanda", "Númer",
                "Kennitala viðtakanda eða greiðanda", "Einkvæmur lykill",
            ])
            sheet.append([
                "29.06.2026", -7069, 1565723, "ISK", "Nova hf.", None,
                "5312050810", "Reikningur", 2, 28, 37, None, None,
                "28.06.2026", "Nova hf.", None, "531205-0810", "13177265277",
            ])
            workbook.save(path)

            bank_type, data = detect_bank_type(path)

        self.assertEqual(bank_type, "arion_reikningsyfirlit")
        self.assertEqual(len(data), 1)
        self.assertEqual(data.loc[0, "Skýring"], "Nova hf.")
        self.assertEqual(
            data.loc[0, "Kennitala viðtakanda eða greiðanda"],
            "531205-0810",
        )

    def test_arion_statement_amount_source_uses_isk_value_for_foreign_currency(self):
        amount_header = "Upph\u00e6\u00f0"
        isk_amount_header = "Upph\u00e6\u00f0 \u00ed ISK"
        frame = pd.DataFrame({
            "Mynt": ["EUR", "ISK"],
            amount_header: [-10, -2000],
            isk_amount_header: [-1500, -2000],
        })

        amounts = _arion_amount_source(frame, amount_header)

        self.assertEqual(amounts.tolist(), [-1500, -2000])


if __name__ == "__main__":
    unittest.main()
